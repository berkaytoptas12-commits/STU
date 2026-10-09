"""Source loaders. Every loader returns a flat list of page-tagged blocks plus the document outline."""

from __future__ import annotations

import html
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional

SUPPORTED_SUFFIXES = {".pdf", ".txt", ".md", ".markdown", ".html", ".htm", ".docx"}

TABLE_KIND = "table"
TEXT_KIND = "text"


@dataclass
class Block:
    page: int  # 1-based
    text: str
    kind: str = TEXT_KIND
    ref: Optional[int] = None  # index into the document's extracted tables (VLM), if any
    bbox: Optional[tuple] = None  # PDF word space (see techrag.geometry), when known
    evidence: bool = True  # False: search-only content (unverified VLM table), never cited directly


@dataclass
class PageStats:
    drawings: int = 0
    images: int = 0
    pdf_tables: int = 0


@dataclass
class LoadedDocument:
    path: Path
    title: str
    blocks: list[Block]
    n_pages: int
    toc: list[tuple[int, str, int]] = field(default_factory=list)  # (level, title, page)
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    page_stats: dict[int, PageStats] = field(default_factory=dict)
    pages: dict = field(default_factory=dict)  # page -> techrag.geometry.PageGeom (PDF text layer words)


def title_from_path(path: Path) -> str:
    return re.sub(r"[_]+", " ", path.stem).strip()


def load_document(path: Path, extract_tables: bool = True) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path, extract_tables)
    if suffix in (".txt", ".md", ".markdown"):
        return load_text(path)
    if suffix in (".html", ".htm"):
        return load_html(path)
    if suffix == ".docx":
        return load_docx(path)
    raise ValueError(f"Unsupported file type: {path}")


# ---------------------------------------------------------------------------- PDF

def _overlap_ratio(a, b) -> float:
    """Fraction of rect a covered by rect b. Rects are (x0, y0, x1, y1)."""
    x0, y0 = max(a[0], b[0]), max(a[1], b[1])
    x1, y1 = min(a[2], b[2]), min(a[3], b[3])
    if x1 <= x0 or y1 <= y0:
        return 0.0
    area = max((a[2] - a[0]) * (a[3] - a[1]), 1e-6)
    return (x1 - x0) * (y1 - y0) / area


def _table_markdown(table) -> str | None:
    try:
        rows = table.extract()
    except Exception:
        return None
    rows = [[_cell(c) for c in row] for row in rows if row and any(c not in (None, "") for c in row)]
    if len(rows) < 2 or max(len(r) for r in rows) < 2:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _cell(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).replace("|", "/").strip()


def _drawing_count(page) -> int:
    try:
        getter = getattr(page, "get_cdrawings", None) or page.get_drawings
        return len(getter())
    except Exception:
        return 0


def load_pdf(path: Path, extract_tables: bool = True) -> LoadedDocument:
    import pymupdf

    from techrag.geometry import PageGeom

    getattr(pymupdf, "no_recommend_layout", lambda: None)()  # silence a stdout advert in newer versions
    doc = pymupdf.open(str(path))
    blocks: list[Block] = []
    warnings: list[str] = []
    empty_pages = 0
    page_stats: dict[int, PageStats] = {}
    geoms: dict[int, PageGeom] = {}
    for index, page in enumerate(doc):
        page_no = index + 1
        try:
            geoms[page_no] = PageGeom.from_page(page, page_no)
        except Exception as exc:  # geometry is for highlighting/grounding only; never fail the document
            warnings.append(f"p.{page_no}: word positions unavailable ({exc.__class__.__name__})")
        tables = []
        stats = page_stats[page_no] = PageStats(drawings=_drawing_count(page))
        try:
            stats.images = len(page.get_images())
        except Exception:
            pass
        # Table detection looks for ruling lines; pages without vector drawings cannot have such tables,
        # and skipping them makes ingestion of long text-only chapters much faster.
        if extract_tables and stats.drawings >= 4:
            try:
                found = page.find_tables()
                for t in found.tables:
                    md = _table_markdown(t)
                    if md:
                        tables.append((tuple(t.bbox), md))
            except Exception as exc:  # table detection is best-effort
                warnings.append(f"p.{page_no}: table extraction failed ({exc.__class__.__name__})")

        items: list[tuple[float, Block]] = []
        for b in page.get_text("blocks", sort=True):
            x0, y0, x1, y1, text, _no, btype = b[:7]
            if btype != 0 or not text.strip():
                continue
            if any(_overlap_ratio((x0, y0, x1, y1), tb) > 0.5 for tb, _ in tables):
                continue
            items.append((y0, Block(page_no, text, bbox=(x0, y0, x1, y1))))
        # Insert each table before the first text block that starts below it (reading order).
        for bbox, md in tables:
            pos = next((i for i, (y, _) in enumerate(items) if y > bbox[1]), len(items))
            items.insert(pos, (bbox[1], Block(page_no, md, TABLE_KIND, bbox=tuple(bbox))))

        stats.pdf_tables = len(tables)
        if not items or sum(len(b.text) for _, b in items) < 20:
            if stats.images:
                empty_pages += 1
        blocks.extend(b for _, b in items)

    if empty_pages:
        warnings.append(
            f"{empty_pages} page(s) contain images but no extractable text - the PDF may be scanned; "
            f"run OCR first (e.g. 'ocrmypdf in.pdf out.pdf')."
        )

    toc = []
    try:
        for entry in doc.get_toc(simple=True):
            level, title, page = int(entry[0]), str(entry[1]).strip(), int(entry[2])
            if title and page >= 1:
                toc.append((level, title, page))
    except Exception:
        pass

    meta = {k: v for k, v in (doc.metadata or {}).items() if v}
    n_pages = doc.page_count
    doc.close()
    return LoadedDocument(path, title_from_path(path), blocks, n_pages, toc, meta, warnings, page_stats, geoms)


# --------------------------------------------------------------------- text / md

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def _paragraph_blocks(text: str, page: int = 1) -> list[Block]:
    out = []
    for para in re.split(r"\n\s*\n", text):
        if para.strip():
            out.append(Block(page, para.strip("\n")))
    return out


def load_text(path: Path) -> LoadedDocument:
    return _text_document(path, path.read_text(encoding="utf-8", errors="replace"))


def _text_document(path: Path, raw: str) -> LoadedDocument:
    # Form feeds mark page breaks in text exports of standards.
    pages = raw.split("\f")
    blocks: list[Block] = []
    for i, page_text in enumerate(pages):
        for b in _paragraph_blocks(page_text, i + 1):
            # Markdown headings must stand alone so the structure pass can see them.
            buf: list[str] = []
            for line in b.text.split("\n"):
                if _MD_HEADING.match(line):
                    if buf:
                        blocks.append(Block(b.page, "\n".join(buf)))
                        buf = []
                    blocks.append(Block(b.page, line.strip()))
                else:
                    buf.append(line)
            if buf:
                blocks.append(Block(b.page, "\n".join(buf)))
    return LoadedDocument(path, title_from_path(path), blocks, len(pages))


# -------------------------------------------------------------------------- HTML

class _HTMLText(HTMLParser):
    BLOCK_TAGS = {"p", "div", "br", "li", "tr", "table", "section", "article", "pre",
                  "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag in self.BLOCK_TAGS:
            self.parts.append("\n\n" if tag != "br" else "\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def load_html(path: Path) -> LoadedDocument:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    text = html.unescape("".join(parser.parts))
    text = re.sub(r"[ \t]+", " ", text)
    return _text_document(path, text)


# -------------------------------------------------------------------------- DOCX

def load_docx(path: Path) -> LoadedDocument:
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError("python-docx is required for .docx sources (pip install python-docx)") from exc
    d = docx.Document(str(path))
    blocks: list[Block] = []
    toc: list[tuple[int, str, int]] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, d)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            m = re.match(r"heading\s*(\d)", style)
            if m:
                toc.append((int(m.group(1)), text, 1))
            blocks.append(Block(1, text))
        elif tag == "tbl":
            rows = [[_cell(c.text) for c in row.cells] for row in Table(child, d).rows]
            if len(rows) >= 2:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                md = "\n".join(
                    ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
                    + ["| " + " | ".join(r) + " |" for r in rows[1:]]
                )
                blocks.append(Block(1, md, TABLE_KIND))
    return LoadedDocument(path, title_from_path(path), blocks, 1, toc)
