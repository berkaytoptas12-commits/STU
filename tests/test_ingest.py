from pathlib import Path

from conftest import HEADER, make_standard_pdf

from techrag.config import Config
from techrag.ingest.chunker import chunk_blocks, estimate_tokens, split_table
from techrag.ingest.cleaning import normalize_text, remove_headers_footers, toc_like_pages
from techrag.ingest.loaders import Block, load_document
from techrag.ingest.pipeline import build_chunks
from techrag.ingest.structure import SectionedBlock, detect_heading


def test_pdf_pipeline_cleans_and_structures(tmp_path: Path):
    pdf = make_standard_pdf(tmp_path / "ddr4.pdf")
    doc, blocks, chunks = build_chunks(Config(), pdf)

    assert doc.n_pages == 7
    assert len(doc.toc) == 5
    text = "\n".join(c.text for c in chunks)
    assert HEADER not in text, "running header must be removed"
    assert "Page 4" not in text, "footer must be removed"
    assert "Table of Contents" not in text, "printed TOC page must be skipped"

    by_section = {c.section[-1]: c for c in chunks if c.section and c.kind == "text"}
    timing = by_section["3.2 Refresh Timing"]
    assert "350 ns" in timing.text
    assert timing.section == ("3 Refresh", "3.2 Refresh Timing")
    assert (timing.page_start, timing.page_end) == (5, 6), "short parent section folded into its first child"
    assert "1.2 V" in by_section["2 Electrical Interface"].text

    tables = [c for c in chunks if c.kind == "table"]
    assert len(tables) == 1
    assert tables[0].text.startswith("Table 3-1 Timing parameters"), "caption is attached to the table"
    assert "| tRFC | 350 | - | ns |" in tables[0].text
    assert doc.page_stats[7].drawings >= 30, "figure page drawings are counted"


def test_normalize_text_dehyphenates_and_fixes_ligatures():
    assert normalize_text("termi-\nnation ﬁeld") == "termination field"
    assert normalize_text("a    b") == "a b"


def test_header_footer_removal_keeps_body():
    blocks = []
    for p in range(1, 8):
        blocks += [Block(p, "PCI Express Base Specification, Rev. 5.0"), Block(p, f"Body text page {p} unique."),
                   Block(p, f"Page {p} of 7")]
    out = remove_headers_footers(blocks, 7)
    texts = [b.text for b in out]
    assert texts == [f"Body text page {p} unique." for p in range(1, 8)]


def test_toc_page_detection():
    lines = "\n".join(f"4.{i} Section title ........ {10 + i}" for i in range(10))
    blocks = [Block(2, lines), Block(3, "Normal paragraph.\nAnother line.\nMore.\nAnd more.\nEnd.")]
    assert toc_like_pages(blocks) == {2}


def test_heading_detection():
    assert detect_heading("4.2.6.3 Polling", True) == (4, "4.2.6.3 Polling")
    assert detect_heading("Appendix B Compliance", True) == (1, "Appendix B Compliance")
    assert detect_heading("3.3 V supply rail", True) is None
    assert detect_heading("2. The device shall transmit", False) is None
    assert detect_heading("1 100 200 300", True) is None


def test_markdown_headings_become_sections(tmp_path: Path):
    md = tmp_path / "notes.md"
    md.write_text("# Top\n\nIntro text here.\n\n## 3.1 Speeds\n\nFast-mode is 400 kbit/s.\n", encoding="utf-8")
    _, _, chunks = build_chunks(Config(), md)
    assert any(c.section == ("Top", "3.1 Speeds") and "400 kbit/s" in c.text for c in chunks)


def test_chunker_respects_sections_and_overlap():
    sent = "The receiver shall tolerate a common mode voltage within the specified range. "
    blocks = [SectionedBlock(1, "1 A", "text", ("1 A",))]
    blocks += [SectionedBlock(1 + i // 3, sent * 6, "text", ("1 A",)) for i in range(9)]
    blocks += [SectionedBlock(9, "2 B\n\n" + sent * 3, "text", ("2 B",))]
    chunks = chunk_blocks(blocks, target_tokens=200, max_tokens=300, overlap_tokens=40, min_tokens=20)
    assert all(c.tokens <= 360 for c in chunks)
    assert {c.section for c in chunks} == {("1 A",), ("2 B",)}
    a_chunks = [c for c in chunks if c.section == ("1 A",)]
    assert len(a_chunks) >= 3
    assert any(c.overlap > 0 for c in a_chunks[1:])
    for c in a_chunks[1:]:
        if c.overlap:
            assert c.text[:c.overlap].strip() in a_chunks[a_chunks.index(c) - 1].text
    assert a_chunks[0].page_start == 1 and a_chunks[-1].page_end == 3


def test_split_table_repeats_header():
    rows = ["| P | V |", "|---|---|"] + [f"| param{i} | {i} ns |" for i in range(200)]
    parts = split_table("\n".join(rows), max_tokens=120)
    assert len(parts) > 1
    assert all(p.startswith("| P | V |\n|---|---|") for p in parts)
    assert all(estimate_tokens(p) <= 140 for p in parts)


def test_text_loader_form_feed_pages(tmp_path: Path):
    f = tmp_path / "doc.txt"
    f.write_text("page one text\fpage two text", encoding="utf-8")
    doc = load_document(f)
    assert [b.page for b in doc.blocks] == [1, 2]


def test_text_files_with_bom_or_turkish_ansi(tmp_path):
    bom = tmp_path / "a.md"
    bom.write_bytes("﻿# Isı Kontrolü\n\nSıcaklık 85 °C.\n".encode("utf-8"))
    doc = load_document(bom)
    assert not doc.blocks[0].text.startswith("﻿") and doc.blocks[0].text.startswith("# Isı")
    ansi = tmp_path / "b.txt"
    ansi.write_bytes("Çalışma sıcaklığı 85 °C'dir.".encode("cp1254"))
    assert "Çalışma sıcaklığı" in load_document(ansi).blocks[0].text


def test_docx_sections_come_from_styled_headings(tmp_path):
    import docx

    from techrag.ingest.structure import assign_sections

    d = docx.Document()
    d.add_heading("2 Timing", level=1)
    d.add_paragraph("Electrical characteristics are given in clause 3.")
    d.add_heading("2.1 Refresh", level=2)
    d.add_paragraph("tRFC is 350 ns.")
    d.add_heading("3 Electrical characteristics", level=1)
    d.add_paragraph("VDD is 1.2 V.")
    path = tmp_path / "spec.docx"
    d.save(str(path))
    doc = load_document(path)
    secs = {b.text: b.section for b in assign_sections(doc, doc.blocks)}
    assert secs["Electrical characteristics are given in clause 3."] == ("2 Timing",)
    assert secs["tRFC is 350 ns."] == ("2 Timing", "2.1 Refresh")
    assert secs["VDD is 1.2 V."] == ("3 Electrical characteristics",)
    assert [t[1] for t in doc.toc] == ["2 Timing", "2.1 Refresh", "3 Electrical characteristics"]


def test_vlm_reply_cut_off_is_reported_not_cached(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import pytest

    from techrag.ingest import vlm
    from techrag.llm import ChatResult

    monkeypatch.setattr(vlm, "render_page", lambda *a, **k: b"png")
    reply = {"r": ChatResult(content='{"tables": [{"caption": "T1", "rows": [["a"', finish_reason="length")}
    client = SimpleNamespace(model="m", chat=lambda *a, **k: reply["r"])
    ex = vlm.TableExtractor(client, tmp_path / "cache")
    with pytest.raises(ValueError, match="cut off"):
        ex.extract_page(tmp_path / "x.pdf", "a" * 64, 3, "text")
    assert not list((tmp_path / "cache").rglob("*.json")), "nothing cached: the next re-index asks again"
    results, errors = ex.extract(tmp_path / "x.pdf", "a" * 64, [3], {3: "text"})
    assert results[3] == [] and "cut off" in errors[0]
    reply["r"] = ChatResult(content='{"tables": [{"caption": "T", "columns": ["a"], "rows": null}]}', finish_reason="stop")
    assert ex.extract_page(tmp_path / "x.pdf", "b" * 64, 1, "text")[0].caption == "T"


def test_fts_query_keeps_turkish_dotted_capital_i():
    from techrag.store import fts_query

    assert fts_query("İletişim İvme") == '"iletişim" OR "ivme"'
